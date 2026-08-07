import os
import sys
import time
import json
import random
import datetime
import requests
import re
import csv
import builtins
from bs4 import BeautifulSoup
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.chrome.service import Service
from webdriver_manager.chrome import ChromeDriverManager
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

BASE = "https://www.screener.in"

CACHE_FILE = "name_cache.json"
FINANCIAL_CACHE_FILE = "financial_cache.json"
LOG_DIR = "logs"
LOG_FILE = os.path.join(LOG_DIR, "screener_financial_analyzer.log")
os.makedirs(LOG_DIR, exist_ok=True)

sys.stdout.reconfigure(line_buffering=True)


def log(*args, **kwargs):
    timestamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    message = " ".join(str(a) for a in args) if args else ""
    line = f"[{timestamp}] {message}"
    with open(LOG_FILE, "a", encoding="utf-8") as f:
        f.write(line + "\n")
    builtins.print(line, **kwargs, flush=True)


print = log


def timer(func):
    def wrapper(*args, **kwargs):
        start_time = time.time()
        print(f"\n{'='*60}")
        print(f"STARTING: {func.__name__}")
        print(f"Time: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        print(f"{'='*60}\n")
        result = func(*args, **kwargs)
        end_time = time.time()
        minutes = int((end_time - start_time) // 60)
        seconds = (end_time - start_time) % 60
        print(f"\n{'='*60}")
        print(f"COMPLETED: {func.__name__}")
        print(f"Total Time: {minutes} minutes {seconds:.2f} seconds")
        print(f"{'='*60}\n")
        return result
    return wrapper


def sanitize_filename(name):
    invalid_chars = r'[<>:"/\\|?*]'
    sanitized = re.sub(invalid_chars, '', name)
    sanitized = sanitized.replace("'", "").replace('"', '')
    sanitized = ' '.join(sanitized.split())
    return sanitized if sanitized else "Company"


def sanitize_sheet_name(name, max_length=31):
    invalid_chars = r'[\[\]:*?/\\]'
    sanitized = re.sub(invalid_chars, ' ', name)
    sanitized = ' '.join(sanitized.split())
    if len(sanitized) > max_length:
        sanitized = sanitized[:max_length-3] + "..."
    return sanitized if sanitized else "Sheet"


def load_cache():
    if not os.path.exists(CACHE_FILE):
        return {}
    with open(CACHE_FILE, "r", encoding="utf-8") as f:
        return json.load(f)


def save_cache(cache):
    with open(CACHE_FILE, "w", encoding="utf-8") as f:
        json.dump(cache, f, indent=2, ensure_ascii=False)


def load_financial_cache():
    if not os.path.exists(FINANCIAL_CACHE_FILE):
        return {}
    with open(FINANCIAL_CACHE_FILE, "r", encoding="utf-8") as f:
        return json.load(f)


def save_financial_cache(cache):
    with open(FINANCIAL_CACHE_FILE, "w", encoding="utf-8") as f:
        json.dump(cache, f, indent=2, ensure_ascii=False)


def to_number(val):
    if not val or val == "-" or val == "":
        return None
    try:
        val_str = str(val).replace(",", "").replace("%", "").strip()
        return float(val_str) if val_str else None
    except:
        return None


def safe_round(n, d=2):
    if n is None or not isinstance(n, (int, float)):
        return None
    try:
        return round(n, d)
    except:
        return None


def extract_company_urls_from_screen(screen_url):
    all_urls = set()
    page = 1
    last_batch = set()

    while True:
        paged = f"{screen_url}?page={page}"
        log(f"  Fetching page {page} for {screen_url}")
        try:
            r = requests.get(paged, headers={"User-Agent": "Mozilla/5.0"}, timeout=10)
            soup = BeautifulSoup(r.text, "html.parser")
            rows = soup.select("tbody tr a[href^='/company/']")
            urls = {BASE + a.get("href") for a in rows if "/company/" in a.get("href", "")}
            if not urls or urls == last_batch:
                break
            all_urls |= urls
            last_batch = urls
            page += 1
            time.sleep(random.uniform(0.6, 1.4))
        except Exception as e:
            print(f"  ⚠ Error: {str(e)}")
            break
    return sorted(all_urls)


def get_company_name(company_url, cache):
    if company_url in cache:
        return cache[company_url]
    try:
        r = requests.get(company_url, headers={"User-Agent": "Mozilla/5.0"}, timeout=10)
        soup = BeautifulSoup(r.text, "html.parser")
        h1 = soup.select_one("h1")
        if h1:
            clean = h1.text.strip()
        else:
            title = soup.title.text if soup.title else ""
            clean = title.split("|")[0].strip()
        cache[company_url] = clean
        time.sleep(random.uniform(0.5, 1.1))
        return clean
    except Exception as e:
        print(f"  ⚠ Error: {str(e)}")
        return company_url.rstrip("/").split("/")[-1]


def setup_driver():
    chrome_options = Options()
    chrome_options.add_argument("--headless")
    chrome_options.add_argument("--no-sandbox")
    chrome_options.add_argument("--disable-dev-shm-usage")
    chrome_options.add_argument("--disable-gpu")
    chrome_options.add_argument("--window-size=1920,1080")
    chrome_options.add_argument("--user-agent=Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36")
    
    service = Service(ChromeDriverManager().install())
    driver = webdriver.Chrome(service=service, options=chrome_options)
    return driver


def expand_all_sections(driver):
    script = """
    async function expandAll() {
        let changed = true;
        let iterations = 0;
        
        while (changed && iterations < 15) {
            changed = false;
            iterations++;
            
            let buttons = document.querySelectorAll('button');
            for (let btn of buttons) {
                if (btn.textContent.includes('+')) {
                    try {
                        btn.scrollIntoView({block: 'center'});
                        await new Promise(r => setTimeout(r, 200));
                        btn.click();
                        await new Promise(r => setTimeout(r, 500));
                        changed = true;
                    } catch(e) {}
                }
            }
            
            await new Promise(r => setTimeout(r, 800));
        }
        
        let yearlyBtn = document.querySelector('button[data-tab-id="yearly-shp"]');
        if (yearlyBtn) {
            yearlyBtn.scrollIntoView({block: 'center'});
            yearlyBtn.click();
            await new Promise(r => setTimeout(r, 1500));
            
            let shpButtons = document.querySelectorAll('#yearly-shp button');
            for (let btn of shpButtons) {
                if (btn.textContent.includes('+')) {
                    try {
                        btn.click();
                        await new Promise(r => setTimeout(r, 400));
                    } catch(e) {}
                }
            }
        }
        
        return iterations;
    }
    return expandAll();
    """
    
    try:
        iterations = driver.execute_script(script)
        print(f"      Expanded in {iterations} iterations")
        time.sleep(2)
    except Exception as e:
        print(f"      ⚠ Expansion error: {str(e)}")


def extract_all_rows_from_table(table):
    rows = []
    all_trs = table.select("tr")
    
    for tr in all_trs:
        label_cell = tr.select_one("td.text")
        if not label_cell:
            continue
        
        button = label_cell.select_one("button")
        if button:
            label = button.get_text(separator=" ", strip=True)
            label = re.sub(r'\s*[−+]\s*$', '', label).strip()
            indent = 0
        else:
            label = label_cell.get_text(separator=" ", strip=True).strip()
            style = label_cell.get("style", "")
            if "padding-left: 2.5rem" in style or "padding-left:2.5rem" in style:
                indent = 1
            elif "padding-left: 5rem" in style or "padding-left:5rem" in style:
                indent = 2
            else:
                classes = tr.get("class", [])
                if "sub" in classes:
                    indent = 1
                elif "nested" in classes:
                    indent = 2
                else:
                    indent = 0
        
        if not label or "Raw PDF" in label or label == "":
            continue
        
        values = []
        for td in tr.select("td:not(.text)"):
            val = td.text.strip()
            values.append(val)
        
        if values:
            rows.append({
                "label": label,
                "values": values,
                "indent": indent
            })
    
    return rows


def extract_table_data(soup, section_id, section_name):
    section = soup.select_one(section_id)
    if not section:
        return None
    
    table = section.select_one("table")
    if not table:
        return None
    
    headers = []
    thead = table.select_one("thead")
    if thead:
        headers = [th.text.strip() for th in thead.select("th") if th.text.strip()]
    
    if not headers:
        return None
    
    rows = extract_all_rows_from_table(table)
    
    if rows:
        return {
            "name": section_name,
            "headers": headers,
            "rows": rows
        }
    return None


def extract_shareholding(soup):
    yearly_div = soup.select_one("#yearly-shp")
    quarterly_div = soup.select_one("#quarterly-shp")
    
    active_div = None
    if yearly_div and "hidden" not in yearly_div.get("class", []):
        active_div = yearly_div
    elif quarterly_div:
        active_div = quarterly_div
    
    if not active_div:
        return None
    
    table = active_div.select_one("table")
    if not table:
        return None
    
    headers = []
    thead = table.select_one("thead")
    if thead:
        headers = [th.text.strip() for th in thead.select("th") if th.text.strip()]
    
    if not headers:
        return None
    
    rows = []
    tbody = table.select_one("tbody")
    if tbody:
        all_trs = tbody.select("tr")
        
        for tr in all_trs:
            label_cell = tr.select_one("td.text")
            if not label_cell:
                continue
            
            button = label_cell.select_one("button")
            if button:
                label = button.get_text(separator=" ", strip=True)
                label = re.sub(r'\s*[−+]\s*$', '', label).strip()
                indent = 0
            else:
                label = label_cell.get_text(separator=" ", strip=True).strip()
                indent = 0
            
            if not label:
                continue
            
            values = []
            for td in tr.select("td:not(.text)"):
                val = td.text.strip()
                if label == "No. of Shareholders":
                    val = val.replace(",", "")
                values.append(val)
            
            if values:
                rows.append({
                    "label": label,
                    "values": values,
                    "indent": indent
                })
    
    name = "Shareholding Pattern (Yearly)" if active_div.get("id") == "yearly-shp" else "Shareholding Pattern (Quarterly)"
    
    return {
        "name": name,
        "headers": headers,
        "rows": rows
    }


def calc_change(values, lag, is_percent):
    res = []
    for i in range(len(values)):
        if i < lag:
            res.append("N/A")
            continue
        
        curr = to_number(values[i])
        prev = to_number(values[i - lag])
        
        if curr is None or prev is None:
            res.append("N/A")
            continue
        
        if is_percent:
            diff = safe_round(curr - prev, 2)
            if diff is not None:
                sign = "+" if diff >= 0 else ""
                res.append(f"{sign}{diff}%")
            else:
                res.append("N/A")
        else:
            diff = safe_round(curr - prev, 0)
            if diff is not None and prev != 0:
                pct = safe_round((diff / prev) * 100, 2)
                sign = "+" if diff >= 0 else ""
                formatted_diff = f"{int(diff):,}" if diff >= 0 else f"{int(diff):,}"
                res.append(f"{sign}{formatted_diff} ({sign}{pct}%)")
            else:
                res.append("N/A")
    return res


def is_percent_metric(label):
    return any(x in label for x in ["%", "Tax %", "OPM %", "ROCE %", "Growth %", "Cost %", "Payout %"]) or label == "CFO/OP"


def extract_all_financials_selenium(company_url):
    log("    Loading company page...")
    
    driver = setup_driver()
    
    try:
        driver.get(company_url)
        time.sleep(4)
        
        log("    Expanding all sections...")
        expand_all_sections(driver)
        
        html = driver.page_source
        soup = BeautifulSoup(html, "html.parser")
        
        sections = []
        
        section_configs = [
            ("#quarters", "Quarterly Results"),
            ("#profit-loss", "Profit & Loss"),
            ("#balance-sheet", "Balance Sheet"),
            ("#cash-flow", "Cash Flows"),
            ("#ratios", "Ratios")
        ]
        
        for section_id, name in section_configs:
            data = extract_table_data(soup, section_id, name)
            if data:
                sections.append(data)
                main = sum(1 for r in data["rows"] if r["indent"] == 0)
                sub = sum(1 for r in data["rows"] if r["indent"] > 0)
                log(f"      ✓ {name}: {len(data['rows'])} metrics ({main} main, {sub} nested)")
        
        sh_data = extract_shareholding(soup)
        if sh_data:
            sections.append(sh_data)
            main = sum(1 for r in sh_data["rows"] if r["indent"] == 0)
            sub = sum(1 for r in sh_data["rows"] if r["indent"] > 0)
            log(f"      ✓ {sh_data['name']}: {len(sh_data['rows'])} metrics ({main} main, {sub} nested)")
        
        return sections if sections else None
        
    except Exception as e:
        print(f"    ⚠ Error: {str(e)}")
        import traceback
        traceback.print_exc()
        return None
    finally:
        driver.quit()


def create_company_csv_file(company_name, financial_data, output_dir):
    safe_filename = sanitize_filename(company_name)
    filename = f"{safe_filename}_Financial_Analysis.csv"
    filepath = os.path.join(output_dir, filename)
    
    with open(filepath, 'w', newline='', encoding='utf-8-sig') as f:
        writer = csv.writer(f)
        
        writer.writerow([f"Financial Analysis - {company_name}"])
        writer.writerow([])
        
        for section in financial_data:
            writer.writerow([section["name"]])
            writer.writerow([])
            
            for row_data in section["rows"]:
                label = row_data["label"]
                values = row_data["values"]
                indent = row_data.get("indent", 0)
                
                display_label = label
                if indent == 1:
                    display_label = "  " + label
                elif indent >= 2:
                    display_label = "    " + label
                
                is_percent = is_percent_metric(label)
                
                headers = section["headers"]
                row_to_write = ["Metric"] + headers
                writer.writerow(row_to_write)
                
                row_to_write = [display_label] + values
                writer.writerow(row_to_write)
                
                if section["name"] == "Quarterly Results":
                    qoq = calc_change(values, 1, is_percent)
                    row_to_write = ["QoQ"] + qoq
                    writer.writerow(row_to_write)
                
                if section["name"] == "Quarterly Results":
                    change_label = "YoY"
                    lag = 4
                else:
                    change_label = "Change"
                    lag = 1
                
                changes = calc_change(values, lag, is_percent)
                row_to_write = [change_label] + changes
                writer.writerow(row_to_write)
                
                writer.writerow([])
            
            writer.writerow([])
    
    return filepath


def create_master_excel(all_companies, subfolders, base_output_dir):
    """Create master Excel file with sheets for each subfolder"""
    
    today = datetime.date.today().strftime('%d_%m_%Y')
    excel_filename = os.path.join(base_output_dir, f"MASTER_ANALYSIS_TRACKER_{today}.xlsx")
    
    wb = openpyxl.Workbook()
    
    # Remove default sheet
    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])
    
    # Style definitions
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=12)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    for folder_name, companies in subfolders.items():
        sheet_name = sanitize_sheet_name(folder_name)
        ws = wb.create_sheet(sheet_name)
        
        # Headers
        headers = ["Company Name", "Classification", "Key Reason 1", "Key Reason 2", "Key Reason 3", "Risk 1", "Risk 2", "Confidence Score (1-10)"]
        
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        # Add companies
        for row_idx, company_name in enumerate(sorted(companies), 2):
            cell = ws.cell(row=row_idx, column=1, value=company_name)
            cell.border = border
            cell.alignment = Alignment(horizontal='left', vertical='center')
            
            # Add empty cells with borders for the remaining columns
            for col_idx in range(2, 9):
                cell = ws.cell(row=row_idx, column=col_idx, value="")
                cell.border = border
                cell.alignment = Alignment(horizontal='center', vertical='center')
        
        # Adjust column widths
        ws.column_dimensions['A'].width = 40
        for col in ['B', 'C', 'D', 'E', 'F', 'G', 'H']:
            ws.column_dimensions[col].width = 20
        
        print(f"  ✓ Sheet '{sheet_name}': {len(companies)} companies")
    
    # Create Summary sheet
    summary_ws = wb.create_sheet("Summary", 0)
    summary_data = []
    for folder_name, companies in subfolders.items():
        summary_data.append({
            "Folder": folder_name,
            "Number of Companies": len(companies)
        })
    summary_data.append({
        "Folder": "TOTAL",
        "Number of Companies": sum(len(c) for c in subfolders.values())
    })
    
    for col_idx, header in enumerate(["Folder", "Number of Companies"], 1):
        cell = summary_ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal='center', vertical='center')
    
    for row_idx, data in enumerate(summary_data, 2):
        cell = summary_ws.cell(row=row_idx, column=1, value=data["Folder"])
        cell.border = border
        cell = summary_ws.cell(row=row_idx, column=2, value=data["Number of Companies"])
        cell.border = border
        cell.alignment = Alignment(horizontal='center', vertical='center')
    
    summary_ws.column_dimensions['A'].width = 30
    summary_ws.column_dimensions['B'].width = 20
    
    wb.save(excel_filename)
    print(f"\n✓ Master Excel created: {excel_filename}")
    
    return excel_filename


@timer
def process_all_screens_and_financials():
    SCREENS = {
        "Early spot": "https://www.screener.in/screens/3603992/early-spot/",
        "Early turnaround": "https://www.screener.in/screens/3604013/early-turnaround/",
        "Strong growth": "https://www.screener.in/screens/3604284/strong-growth/",
        "Early Turnarounds": "https://www.screener.in/screens/3608038/early-turnarounds/",
        "Stable Turnarounds": "https://www.screener.in/screens/3608044/stable-turnarounds/"
    }
    
    COMPANIES_PER_FOLDER = 10
    
    log("\n" + "="*60)
    log("PHASE 1: EXTRACTING UNIQUE COMPANIES")
    log("="*60)
    
    all_companies = {}
    name_cache = load_cache()
    
    for screen_name, screen_url in SCREENS.items():
        log(f"\n{'─'*50}")
        log(f"Processing screen: {screen_name}")
        log(f"{'─'*50}")
        
        company_urls = extract_company_urls_from_screen(screen_url)
        log(f"  Found {len(company_urls)} companies")
        
        for url in company_urls:
            company_name = get_company_name(url, name_cache)
            if company_name not in all_companies:
                all_companies[company_name] = {"url": url}
    
    save_cache(name_cache)
    log(f"\n{'='*60}")
    log(f"Total unique companies: {len(all_companies)}")
    log(f"{'='*60}")
    
    # Save JSON
    companies_json = {name: data["url"] for name, data in all_companies.items()}
    json_filename = f"unique_companies_{datetime.date.today().strftime('%d_%m_%Y')}.json"
    with open(json_filename, "w", encoding="utf-8") as f:
        json.dump(companies_json, f, indent=2, ensure_ascii=False)
    log(f"✓ Saved: {json_filename}")
    
    log("="*60)
    log("PHASE 2: EXTRACTING FINANCIAL DATA")
    log("="*60)
    
    today = datetime.date.today().strftime('%d_%m_%Y')
    base_output_dir = f"Financial_Analysis_{today}"
    os.makedirs(base_output_dir, exist_ok=True)
    log(f"✓ Base output directory: {base_output_dir}")
    
    # Create subfolders and distribute companies
    sorted_companies = sorted(all_companies.items())
    subfolders = {}
    folder_companies = {}
    
    for i in range(0, len(sorted_companies), COMPANIES_PER_FOLDER):
        folder_num = i // COMPANIES_PER_FOLDER + 1
        folder_name = f"Batch_{folder_num:03d}"
        folder_path = os.path.join(base_output_dir, folder_name)
        os.makedirs(folder_path, exist_ok=True)
        
        batch = sorted_companies[i:i+COMPANIES_PER_FOLDER]
        subfolders[folder_name] = folder_path
        folder_companies[folder_name] = [name for name, _ in batch]
        
        log(f"✓ Created: {folder_name} ({len(batch)} companies)")
    
    log(f"\n✓ Total {len(subfolders)} subfolders created")
    
    financial_cache = load_financial_cache()
    successful = 0
    failed = 0
    skipped = 0
    
    # Process each company and save to its subfolder
    for folder_name, folder_path in subfolders.items():
        companies_in_folder = folder_companies[folder_name]
        log(f"\n{'─'*50}")
        log(f"Processing {folder_name}")
        log(f"{'─'*50}")
        
        for company_name in companies_in_folder:
            company_url = all_companies[company_name]["url"]
            
            safe_filename = sanitize_filename(company_name)
            csv_path = os.path.join(folder_path, f"{safe_filename}_Financial_Analysis.csv")
            
            if os.path.exists(csv_path):
                print(f"  ⏭ {company_name} - Already exists")
                skipped += 1
                continue
            
            log(f"  Processing: {company_name}")
            
            if company_url in financial_cache:
                print(f"    ✓ Using cached data")
                financial_data = financial_cache[company_url]
            else:
                financial_data = extract_all_financials_selenium(company_url)
                if financial_data:
                    financial_cache[company_url] = financial_data
                    if successful % 5 == 0:
                        save_financial_cache(financial_cache)
                        log("    💾 Cache saved")
                time.sleep(random.uniform(1, 2))
            
            if financial_data:
                try:
                    filepath = create_company_csv_file(company_name, financial_data, folder_path)
                    log(f"    ✓ Saved to {folder_name}")
                    successful += 1
                except Exception as e:
                    print(f"    ⚠ Error: {str(e)}")
                    failed += 1
            else:
                print(f"    ⚠ No data")
                failed += 1
    
    save_financial_cache(financial_cache)
    
    # Create master Excel file
    log("\n" + "="*60)
    log("PHASE 3: CREATING MASTER EXCEL TRACKER")
    log("="*60)
    
    create_master_excel(all_companies, folder_companies, base_output_dir)
    
    log(f"\n{'='*60}")
    log(f"✓ COMPLETE!")
    log(f"✓ Output directory: {base_output_dir}")
    log(f"✓ Subfolders: {len(subfolders)}")
    log(f"✓ CSV files created: {successful}")
    log(f"✓ Failed: {failed}")
    log(f"✓ Skipped: {skipped}")
    log(f"{'='*60}")
    
    return base_output_dir


if __name__ == "__main__":
    try:
        from webdriver_manager.chrome import ChromeDriverManager
    except ImportError:
        os.system("pip install webdriver-manager")
        from webdriver_manager.chrome import ChromeDriverManager
    
    process_all_screens_and_financials()